#!/usr/bin/env python3
"""
==============================================================================
Get Knowledge — Local & Production Backend Server (server.py)
==============================================================================
Features:
1. Serves static website files (HTML, CSS, JS, Assets) on port 8000
2. Handles POST /api/apply:
   - Records application into applications.xlsx
   - Compares incoming UPI Reference (UTR) with existing entries in applications.xlsx
   - Informs the user if a duplicate UPI reference is detected
   - Sends an automatic confirmation email from getknowledge2026@gmail.com
     stating: "greeting + once manual verification done our team contact you
     and you will be added to whatsapp group"
3. Handles GET /api/download-xlsx to download the live applications.xlsx file
4. Handles GET /api/applications to view recorded entries in JSON format
==============================================================================
"""

import os
import sys
import json
import base64
import hashlib
import smtplib
import datetime
import time
import threading
import urllib.request
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import socket
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PORT = int(os.environ.get("PORT", 8000))
XLSX_FILENAME = "applications.xlsx"
UPLOADS_DIR = "uploads"
OFFICIAL_EMAIL = "getknowledge2026@gmail.com"
OFFICIAL_PHONE = "+91 96262 61123"
SMTP_SERVER = os.environ.get("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.environ.get("SMTP_PORT", 587))
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD", "") # Optional: Gmail App Password

def get_smtp_password():
    """Reads Gmail App Password from environment variable, .env, or email_config.json."""
    pwd = os.environ.get("SMTP_PASSWORD", "").strip()
    if pwd:
        return pwd
    # Try .env file
    if os.path.exists(".env"):
        try:
            with open(".env", "r") as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("SMTP_PASSWORD="):
                        return line.split("=", 1)[1].strip().strip('"').strip("'")
        except Exception:
            pass
    # Try email_config.json
    if os.path.exists("email_config.json"):
        try:
            with open("email_config.json", "r") as f:
                cfg = json.load(f)
                return cfg.get("smtpPassword", "").strip()
        except Exception:
            pass
    return ""

def get_admin_password():
    """Retrieve ADMIN_PASSWORD from environment or .env, defaulting to '09876543211'."""
    pwd = os.environ.get("ADMIN_PASSWORD", "").strip()
    if pwd:
        return pwd
    if os.path.exists(".env"):
        try:
            with open(".env", "r") as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("ADMIN_PASSWORD="):
                        return line.split("=", 1)[1].strip().strip('"').strip("'")
        except Exception:
            pass
    return "09876543211"

def get_admin_token():
    """Deterministic secure session token derived from current ADMIN_PASSWORD."""
    pw = get_admin_password()
    return hashlib.sha256(f"gk_portal_salt_2026_{pw}".encode("utf-8")).hexdigest()

def ensure_xlsx_exists():
    """Ensure applications.xlsx exists with styled headers."""
    if os.path.exists(XLSX_FILENAME):
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications"

    headers = [
        "Timestamp",
        "Applicant ID",
        "Full Name",
        "Phone Number",
        "WhatsApp Number",
        "Email Address",
        "Course",
        "Fee (INR)",
        "Batch Timing",
        "Payment Status",
        "UPI Reference / UTR",
        "Screenshot File",
        "Verification Status"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0C2340", end_color="0C2340", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D2D6DC"),
        right=Side(style="thin", color="D2D6DC"),
        top=Side(style="thin", color="D2D6DC"),
        bottom=Side(style="thin", color="D2D6DC")
    )
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.freeze_panes = "A2"

    col_widths = {
        "A": 20, "B": 16, "C": 22, "D": 18, "E": 18, "F": 28,
        "G": 20, "H": 12, "I": 32, "J": 24, "K": 24, "L": 24, "M": 22
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    wb.save(XLSX_FILENAME)
    print(f"[*] Initialized {XLSX_FILENAME}")

def check_duplicate_utr(new_utr):
    """
    Check if the given UPI Reference Number (UTR) already exists in applications.xlsx.
    Returns (is_duplicate: bool, matching_row: dict or None)
    """
    if not new_utr or new_utr.strip().upper() in ["N/A", "NONE", "PENDING", ""]:
        return False, None

    clean_utr = new_utr.strip().upper()
    ensure_xlsx_exists()

    try:
        wb = openpyxl.load_workbook(XLSX_FILENAME, data_only=True)
        ws = wb.active
        # Column K (index 11) is 'UPI Reference / UTR'
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if len(row) >= 11 and row[10]:
                existing_utr = str(row[10]).strip().upper()
                if existing_utr == clean_utr:
                    match_info = {
                        "row": row_idx,
                        "applicantId": row[1] if len(row) > 1 else "",
                        "name": row[2] if len(row) > 2 else "",
                        "timestamp": row[0] if len(row) > 0 else ""
                    }
                    return True, match_info
        return False, None
    except Exception as e:
        print(f"[!] Error checking duplicate UTR: {e}")
        return False, None

def save_base64_screenshot(applicant_id, filename, base64_str):
    """Saves a base64 encoded image to the uploads folder and returns the saved relative file path."""
    if not base64_str:
        return filename or "None Attached"
    try:
        os.makedirs(UPLOADS_DIR, exist_ok=True)
        if "," in base64_str:
            header, base64_data = base64_str.split(",", 1)
        else:
            base64_data = base64_str
            header = ""

        ext = ".png"
        if "jpeg" in header or "jpg" in header or (filename and filename.lower().endswith((".jpg", ".jpeg"))):
            ext = ".jpg"
        elif "webp" in header or (filename and filename.lower().endswith(".webp")):
            ext = ".webp"

        safe_name = f"{applicant_id}_screenshot{ext}"
        filepath = os.path.join(UPLOADS_DIR, safe_name)
        with open(filepath, "wb") as f:
            f.write(base64.b64decode(base64_data))
        print(f"[+] Saved screenshot proof to {filepath}")
        return filepath
    except Exception as e:
        print(f"[!] Failed to save screenshot image: {e}")
        return filename or "None Attached"

def append_to_xlsx(record):
    """Appends a new application record to applications.xlsx."""
    ensure_xlsx_exists()
    wb = openpyxl.load_workbook(XLSX_FILENAME)
    ws = wb.active

    screenshot_val = record.get("screenshotFileName", "None Attached")
    if record.get("screenshotBase64"):
        saved_path = save_base64_screenshot(
            record.get("applicantId", "GK000000"),
            screenshot_val,
            record.get("screenshotBase64")
        )
        screenshot_val = saved_path

    row_data = [
        record.get("timestamp", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        record.get("applicantId", ""),
        record.get("name", ""),
        record.get("phone", ""),
        record.get("whatsapp", record.get("phone", "")),
        record.get("email", ""),
        record.get("course", "AI Development"),
        record.get("fee", 199),
        record.get("batch", ""),
        record.get("paymentStatus", ""),
        record.get("utr", "N/A"),
        screenshot_val,
        "Pending Manual Verification"
    ]
    ws.append(row_data)

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    last_row = ws.max_row
    ws.row_dimensions[last_row].height = 22
    for cell in ws[last_row]:
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

    wb.save(XLSX_FILENAME)
    print(f"[+] Appended applicant {record.get('applicantId')} to {XLSX_FILENAME}")

def update_xlsx_row(applicant_id, fields):
    """Updates an existing applicant row in applications.xlsx."""
    ensure_xlsx_exists()
    wb = openpyxl.load_workbook(XLSX_FILENAME)
    ws = wb.active

    # Column mappings:
    # 0: Timestamp, 1: Applicant ID, 2: Full Name, 3: Phone Number, 4: WhatsApp Number,
    # 5: Email Address, 6: Course, 7: Fee (INR), 8: Batch Timing, 9: Payment Status,
    # 10: UPI Reference / UTR, 11: Screenshot File, 12: Verification Status
    get_val = lambda *keys: next((fields[k] for k in keys if k in fields and fields[k] is not None), None)

    updated = False
    for row in ws.iter_rows(min_row=2):
        if len(row) >= 2 and str(row[1].value).strip() == str(applicant_id).strip():
            v_name = get_val("name", "Full Name")
            if v_name is not None:
                row[2].value = str(v_name).strip()

            v_phone = get_val("phone", "Phone Number")
            if v_phone is not None:
                row[3].value = str(v_phone).strip()

            v_whatsapp = get_val("whatsapp", "WhatsApp Number")
            if v_whatsapp is not None:
                row[4].value = str(v_whatsapp).strip()

            v_email = get_val("email", "Email Address")
            if v_email is not None:
                row[5].value = str(v_email).strip()

            v_course = get_val("course", "Course")
            if v_course is not None:
                row[6].value = str(v_course).strip()

            v_fee = get_val("fee", "Fee (INR)")
            if v_fee is not None:
                try:
                    row[7].value = float(v_fee)
                except Exception:
                    row[7].value = v_fee

            v_batch = get_val("batch", "Batch Timing")
            if v_batch is not None:
                row[8].value = str(v_batch).strip()

            v_pay_status = get_val("paymentStatus", "Payment Status")
            if v_pay_status is not None:
                row[9].value = str(v_pay_status).strip()

            v_utr = get_val("utr", "UPI Reference / UTR")
            if v_utr is not None:
                row[10].value = str(v_utr).strip()

            if "screenshotBase64" in fields and fields["screenshotBase64"]:
                saved_path = save_base64_screenshot(
                    applicant_id,
                    fields.get("screenshotFileName", "screenshot.png"),
                    fields["screenshotBase64"]
                )
                row[11].value = saved_path
            else:
                v_screenshot = get_val("screenshotFileName", "Screenshot File")
                if v_screenshot is not None:
                    row[11].value = str(v_screenshot).strip()

            v_status = get_val("verificationStatus", "Verification Status")
            if v_status is not None:
                row[12].value = str(v_status).strip()

            updated = True
            break

    if updated:
        wb.save(XLSX_FILENAME)
        return True, "Application row updated in Excel file."
    return False, f"Applicant ID '{applicant_id}' was not found in applications.xlsx."

def delete_xlsx_row(applicant_id):
    """Deletes an applicant row from applications.xlsx."""
    ensure_xlsx_exists()
    wb = openpyxl.load_workbook(XLSX_FILENAME)
    ws = wb.active

    target_idx = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if len(row) >= 2 and str(row[1].value).strip() == str(applicant_id).strip():
            target_idx = row_idx
            break

    if target_idx is not None:
        ws.delete_rows(target_idx, 1)
        wb.save(XLSX_FILENAME)
        return True, f"Applicant '{applicant_id}' deleted from applications.xlsx."
    return False, f"Applicant ID '{applicant_id}' was not found in applications.xlsx."

def send_confirmation_email(record):
    """
    Sends automatic greeting and verification confirmation email from getknowledge2026@gmail.com
    to the applicant.
    """
    to_email = record.get("email")
    if not to_email:
        return False, "No email address provided"

    name = record.get("name", "Student")
    course = record.get("course", "AI Development")
    applicant_id = record.get("applicantId", "N/A")
    whatsapp = record.get("whatsapp", record.get("phone", "your phone"))
    utr = record.get("utr", "N/A")

    subject = f"Welcome to Get Knowledge — {course} Application Confirmation ({applicant_id})"

    body_text = f"""Dear {name},

Greetings from Get Knowledge!

Thank you for applying for our 20-Day Live Cohort in {course}. We have received your application successfully.

--- Application Details ---
• Application ID: {applicant_id}
• Course: {course} (20 Days Live Cohort)
• Fee: ₹199 (All-Inclusive)
• Registered Email: {to_email}
• WhatsApp Number: {whatsapp}
• UPI Reference / UTR: {utr}

--- What Happens Next? ---
Our accounts team will perform manual bank verification of your UPI transaction reference ({utr}).
Once manual verification is done, our team will contact you directly on WhatsApp (+91 {whatsapp}) and you will be added to the batch WhatsApp group with your private Google Meet class links and onboarding kit.

If you have any questions or need immediate assistance, please feel free to reply to this email or reach us directly:
• Email: {OFFICIAL_EMAIL}
• Phone / WhatsApp: {OFFICIAL_PHONE}

We look forward to having you in the live cohort!

Warm regards,
Admissions Team
Get Knowledge
Learn More. Achieve More.
{OFFICIAL_EMAIL} | {OFFICIAL_PHONE}
"""

    # 1. Rich HTML Template
    html_content = f"""
    <div style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; max-width: 600px; margin: 0 auto; border: 1px solid #e2e8f0; border-radius: 12px; overflow: hidden; background: #ffffff;">
        <div style="background: #0c2340; color: #ffffff; padding: 26px 20px; text-align: center;">
            <h2 style="margin: 0; font-size: 22px; font-weight: 800; letter-spacing: 0.5px;">Get Knowledge</h2>
            <p style="margin: 6px 0 0 0; color: #38bdf8; font-size: 12px; font-weight: 700; letter-spacing: 0.1em; text-transform: uppercase;">LEARN MORE. ACHIEVE MORE.</p>
        </div>
        <div style="padding: 24px; color: #334155; line-height: 1.6;">
            <p style="font-size: 16px; color: #0f172a; margin-top: 0;">Dear <strong>{name}</strong>,</p>
            <p style="font-size: 15px; color: #0f172a;"><strong>Greetings from Get Knowledge!</strong></p>
            <p>Thank you for applying for our <strong>20-Day Live Cohort in {course}</strong>. Your registration reference is <strong>{applicant_id}</strong> and has been recorded in our official admissions file.</p>
            
            <div style="background: #f8fafc; border: 1.5px solid #e2e8f0; border-radius: 10px; padding: 16px; margin: 20px 0;">
                <h4 style="margin: 0 0 12px 0; color: #0c2340; font-size: 14px; text-transform: uppercase; letter-spacing: 0.05em;">Application Summary</h4>
                <table style="width: 100%; font-size: 13.5px; border-collapse: collapse;">
                    <tr><td style="padding: 5px 0; color: #64748b;">Applicant ID:</td><td style="font-weight: 700; color: #0c2340;">{applicant_id}</td></tr>
                    <tr><td style="padding: 5px 0; color: #64748b;">Course:</td><td style="font-weight: 600;">{course} (20 Days Live)</td></tr>
                    <tr><td style="padding: 5px 0; color: #64748b;">Course Fee:</td><td style="font-weight: 700; color: #0284c7;">₹199 (All-Inclusive)</td></tr>
                    <tr><td style="padding: 5px 0; color: #64748b;">WhatsApp Number:</td><td style="font-weight: 600;">{whatsapp}</td></tr>
                    <tr><td style="padding: 5px 0; color: #64748b;">UPI Ref / UTR:</td><td style="font-weight: 700; color: #0f172a; background: #f1f5f9; padding: 2px 6px; border-radius: 4px;">{utr}</td></tr>
                </table>
            </div>

            <div style="background: #f0fdf4; border-left: 4px solid #10b981; padding: 14px 18px; border-radius: 6px; margin-bottom: 22px;">
                <strong style="color: #065f46; display: block; font-size: 14px; margin-bottom: 4px;">&#10003; What Happens Next?</strong>
                <p style="margin: 0; font-size: 13.5px; color: #166534; line-height: 1.5;">
                    Our accounts desk will manually verify your UPI reference ({utr}) against our bank statement.<br><br>
                    <strong>Once manual verification is done, our team will contact you directly on WhatsApp (+91 {whatsapp}) and you will be added to the batch WhatsApp group</strong> with your private Google Meet class links and onboarding kit!
                </p>
            </div>

            <p style="font-size: 13px; color: #64748b; margin-top: 24px; line-height: 1.5;">
                Have questions? Contact Admissions directly:<br>
                • Email: <a href="mailto:{OFFICIAL_EMAIL}" style="color: #0284c7; text-decoration: none;">{OFFICIAL_EMAIL}</a><br>
                • Phone / WhatsApp: <a href="https://wa.me/919626261123" style="color: #0284c7; text-decoration: none;">{OFFICIAL_PHONE}</a>
            </p>
        </div>
        <div style="background: #f8fafc; padding: 14px; text-align: center; font-size: 12px; color: #94a3b8; border-top: 1px solid #e2e8f0;">
            &copy; 2026 Get Knowledge &bull; Bengaluru, India
        </div>
    </div>
    """

    # 2. Check for HTTPS Email Webhook URL (Bypasses Render's port 587 block!)
    webhook_url = os.environ.get("EMAIL_WEBHOOK_URL", "").strip()
    if not webhook_url and os.path.exists(".env"):
        try:
            with open(".env", "r") as f:
                for line in f:
                    if line.strip().startswith("EMAIL_WEBHOOK_URL="):
                        webhook_url = line.split("=", 1)[1].strip().strip('"').strip("'")
        except Exception:
            pass

    if webhook_url:
        try:
            payload = json.dumps({
                "to": to_email,
                "subject": subject,
                "body": body_text,
                "html": html_content
            }).encode("utf-8")
            req = urllib.request.Request(
                webhook_url,
                data=payload,
                headers={"Content-Type": "application/json"}
            )
            with urllib.request.urlopen(req, timeout=12) as resp:
                res_body = resp.read().decode("utf-8")
                print(f"[+] Sent confirmation email to {to_email} via HTTPS Webhook: {res_body}")
                return True, "Email sent via Webhook"
        except Exception as e:
            print(f"[!] Webhook email dispatch error ({e}). Trying SMTP fallback...")

    # 3. Fallback: Direct Gmail SMTP
    smtp_pass = get_smtp_password()
    if smtp_pass:
        smtp_clean = smtp_pass.replace(" ", "")
        try:
            msg = MIMEMultipart("alternative")
            msg["From"] = f"Get Knowledge Admissions <{OFFICIAL_EMAIL}>"
            msg["To"] = to_email
            msg["Subject"] = subject
            msg["Reply-To"] = OFFICIAL_EMAIL

            # Attach plain-text and HTML
            msg.attach(MIMEText(body_text, "plain"))
            msg.attach(MIMEText(html_content, "html"))

            with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=8) as server:
                server.starttls()
                server.login(OFFICIAL_EMAIL, smtp_clean)
                server.send_message(msg)
            print(f"[+] Sent confirmation email to {to_email} via SMTP")
            return True, "Email sent successfully"
        except Exception as e:
            print(f"[!] SMTP sending failed ({e}). Email logged to server console.")
            return False, str(e)
    else:
        print(f"[*] (No SMTP password or Webhook set) Confirmation email drafted for {to_email}:\n{'-'*60}\n{body_text}\n{'-'*60}")
        return False, "Neither EMAIL_WEBHOOK_URL nor SMTP_PASSWORD set"

class GetKnowledgeRequestHandler(SimpleHTTPRequestHandler):
    """Custom request handler that serves files and handles /api endpoints."""
    protocol_version = "HTTP/1.1"

    def handle(self):
        """Silently handle client socket aborts without blocking or traceback."""
        try:
            super().handle()
        except (BrokenPipeError, ConnectionResetError):
            self.close_connection = True

    def copyfile(self, source, outputfile):
        """Silently handle client socket disconnects during file streaming."""
        try:
            super().copyfile(source, outputfile)
        except (BrokenPipeError, ConnectionResetError):
            pass

    def check_admin_auth(self):
        """Verify request contains valid X-Admin-Token header or ?token= query parameter."""
        expected = get_admin_token()
        token = self.headers.get("X-Admin-Token", "").strip()
        if not token and "?" in self.path:
            query_str = self.path.split("?", 1)[1]
            for part in query_str.split("&"):
                if part.startswith("token="):
                    token = part.split("=", 1)[1].strip()
                    break
        return token == expected

    def do_OPTIONS(self):
        """Handle CORS pre-flight requests."""
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, X-Admin-Token")
        self.end_headers()

    def do_POST(self):
        if self.path == "/api/admin-auth":
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")
            try:
                data = json.loads(post_body)
            except Exception:
                data = {}
            entered_pw = str(data.get("password", "")).strip()
            if entered_pw == get_admin_password():
                self.send_json_response(200, {
                    "status": "success",
                    "token": get_admin_token(),
                    "message": "Access granted"
                })
            else:
                self.send_json_response(401, {
                    "status": "error",
                    "message": "Incorrect password. Access denied."
                })
            return

        if self.path == "/api/apply":
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")

            try:
                data = json.loads(post_body)
            except Exception as e:
                self.send_json_response(400, {
                    "status": "error",
                    "message": f"Invalid JSON payload: {e}"
                })
                return

            new_utr = data.get("utr", "").strip()
            screenshot = data.get("screenshotFileName", "").strip()

            # Required field validation: UPI Reference & Screenshot
            if not new_utr or new_utr.upper() in ["N/A", "NONE", ""]:
                self.send_json_response(400, {
                    "status": "error",
                    "message": "UPI Reference Number (UTR) is required."
                })
                return

            if not screenshot or screenshot.lower() in ["none", "none attached", ""]:
                self.send_json_response(400, {
                    "status": "error",
                    "message": "Payment screenshot image is required."
                })
                return

            # 1. Compare UPI reference number with existing entries in applications.xlsx
            is_duplicate, match_info = check_duplicate_utr(new_utr)
            if is_duplicate:
                print(f"[!] Duplicate UPI Reference detected: {new_utr} (Row {match_info.get('row')})")
                self.send_json_response(409, {
                    "status": "duplicate",
                    "message": f"This UPI Reference Number ({new_utr}) has already been registered in our records for applicant {match_info.get('applicantId')}. If you believe this is an error, please contact our admissions team at {OFFICIAL_EMAIL} or WhatsApp {OFFICIAL_PHONE}."
                })
                return

            # 2. If nothing matches, append row to applications.xlsx
            try:
                append_to_xlsx(data)
            except Exception as e:
                print(f"[!] Error writing to XLSX: {e}")
                self.send_json_response(500, {
                    "status": "error",
                    "message": f"Failed to record application to Excel: {e}"
                })
                return

            # 3. Send automatic confirmation email in background thread so client response is instant!
            try:
                email_thread = threading.Thread(target=send_confirmation_email, args=(dict(data),), daemon=True)
                email_thread.start()
            except Exception as e:
                print(f"[!] Email thread start failed: {e}")

            # 4. Return success response immediately
            self.send_json_response(200, {
                "status": "success",
                "message": "Application recorded successfully in applications.xlsx!",
                "applicantId": data.get("applicantId")
            })
            return

        if self.path == "/api/update-row":
            if not self.check_admin_auth():
                self.send_json_response(401, {"status": "unauthorized", "message": "Admin authorization required."})
                return
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")
            try:
                data = json.loads(post_body)
            except Exception as e:
                self.send_json_response(400, {"status": "error", "message": f"Invalid JSON: {e}"})
                return

            applicant_id = data.get("applicantId")
            if not applicant_id:
                self.send_json_response(400, {"status": "error", "message": "Applicant ID is required."})
                return

            success, msg = update_xlsx_row(applicant_id, data.get("fields", data))
            code = 200 if success else 404
            self.send_json_response(code, {"status": "success" if success else "error", "message": msg})
            return

        if self.path == "/api/delete-row":
            if not self.check_admin_auth():
                self.send_json_response(401, {"status": "unauthorized", "message": "Admin authorization required."})
                return
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")
            try:
                data = json.loads(post_body)
            except Exception as e:
                self.send_json_response(400, {"status": "error", "message": f"Invalid JSON: {e}"})
                return

            applicant_id = data.get("applicantId")
            if not applicant_id:
                self.send_json_response(400, {"status": "error", "message": "Applicant ID is required."})
                return

            success, msg = delete_xlsx_row(applicant_id)
            code = 200 if success else 404
            self.send_json_response(code, {"status": "success" if success else "error", "message": msg})
            return

        if self.path == "/api/add-row":
            if not self.check_admin_auth():
                self.send_json_response(401, {"status": "unauthorized", "message": "Admin authorization required."})
                return
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")
            try:
                data = json.loads(post_body)
            except Exception as e:
                self.send_json_response(400, {"status": "error", "message": f"Invalid JSON: {e}"})
                return

            if not data.get("applicantId"):
                import random
                data["applicantId"] = f"GK{random.randint(100000, 999999)}"

            try:
                append_to_xlsx(data)
                self.send_json_response(200, {
                    "status": "success",
                    "message": f"Added applicant {data['applicantId']} to applications.xlsx",
                    "applicantId": data["applicantId"]
                })
            except Exception as e:
                self.send_json_response(500, {"status": "error", "message": str(e)})
            return

        if self.path == "/api/upload-screenshot":
            if not self.check_admin_auth():
                self.send_json_response(401, {"status": "unauthorized", "message": "Admin authorization required."})
                return
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length).decode("utf-8")
            try:
                data = json.loads(post_body)
            except Exception as e:
                self.send_json_response(400, {"status": "error", "message": f"Invalid JSON: {e}"})
                return

            applicant_id = data.get("applicantId")
            b64_data = data.get("screenshotBase64")
            filename = data.get("fileName", "screenshot.png")

            if not applicant_id or not b64_data:
                self.send_json_response(400, {"status": "error", "message": "Applicant ID and screenshotBase64 are required."})
                return

            saved_path = save_base64_screenshot(applicant_id, filename, b64_data)
            success, msg = update_xlsx_row(applicant_id, {"screenshotFileName": saved_path})
            if success:
                self.send_json_response(200, {
                    "status": "success",
                    "message": f"Screenshot uploaded and saved for {applicant_id}!",
                    "screenshotFile": saved_path,
                    "url": "/" + saved_path
                })
            else:
                self.send_json_response(404, {"status": "error", "message": msg})
            return

        # Default fallback for other POST requests
        self.send_json_response(404, {"status": "error", "message": "Not Found"})

    def do_GET(self):
        if self.path in ["/api/health", "/health", "/ping"]:
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", "16")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(b'{"status":"ok"}\n')
            return

        if self.path.startswith("/api/download-xlsx"):
            if not self.check_admin_auth():
                self.send_response(401)
                self.send_header("Content-Type", "text/plain")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(b"Unauthorized: Admin password required.")
                return
            if not os.path.exists(XLSX_FILENAME):
                ensure_xlsx_exists()
            try:
                with open(XLSX_FILENAME, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f"attachment; filename=\"{XLSX_FILENAME}\"")
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)
                return
            except Exception as e:
                self.send_json_response(500, {"status": "error", "message": str(e)})
                return

        if self.path == "/api/applications":
            if not self.check_admin_auth():
                self.send_json_response(401, {"status": "unauthorized", "message": "Admin authorization required."})
                return
            ensure_xlsx_exists()
            try:
                wb = openpyxl.load_workbook(XLSX_FILENAME, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) <= 1:
                    self.send_json_response(200, {"count": 0, "headers": [], "applications": []})
                    return
                headers = [str(h) for h in rows[0]]
                apps = []
                for row_idx, row in enumerate(rows[1:], start=2):
                    record = {"_rowIndex": row_idx}
                    for i, h in enumerate(headers):
                        val = row[i] if i < len(row) and row[i] is not None else ""
                        record[h] = str(val)

                    s_val = record.get("Screenshot File", "").strip()
                    has_img = False
                    img_url = ""
                    if s_val and s_val.lower() not in ["none", "none attached", "n/a", ""]:
                        if os.path.exists(s_val):
                            has_img = True
                            img_url = "/" + s_val.lstrip("/")
                        elif os.path.exists(os.path.join(UPLOADS_DIR, s_val)):
                            has_img = True
                            img_url = "/" + os.path.join(UPLOADS_DIR, s_val).lstrip("/")
                    record["_hasScreenshot"] = has_img
                    record["_screenshotUrl"] = img_url
                    apps.append(record)
                self.send_json_response(200, {"count": len(apps), "headers": headers, "applications": apps})
                return
            except Exception as e:
                self.send_json_response(500, {"status": "error", "message": str(e)})
                return

        # Clean URL and Course Redirection Handler
        clean_path = self.path.split('?')[0].rstrip('/')
        if clean_path in ["/admin", "/portal", "/dashboard"]:
            self.send_response(302)
            self.send_header("Location", "/admin.html")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return

        if clean_path in ["/course", "/courses", "/courses.html", "/course-details"]:
            self.send_response(302)
            self.send_header("Location", "/course-details.html")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return

        if not os.path.exists("." + clean_path) and os.path.exists("." + clean_path + ".html"):
            self.path = clean_path + ".html" + (("?" + self.path.split('?', 1)[1]) if '?' in self.path else "")

        try:
            return super().do_GET()
        except (BrokenPipeError, ConnectionResetError):
            pass

    def do_HEAD(self):
        clean_path = self.path.split('?')[0].rstrip('/')
        if clean_path in ["/admin", "/portal", "/dashboard"]:
            self.send_response(302)
            self.send_header("Location", "/admin.html")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return

        if clean_path in ["/course", "/courses", "/courses.html", "/course-details"]:
            self.send_response(302)
            self.send_header("Location", "/course-details.html")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return

        if not os.path.exists("." + clean_path) and os.path.exists("." + clean_path + ".html"):
            self.path = clean_path + ".html" + (("?" + self.path.split('?', 1)[1]) if '?' in self.path else "")

        try:
            return super().do_HEAD()
        except (BrokenPipeError, ConnectionResetError):
            pass

    def end_headers(self):
        """Disable browser caching for all static and dynamic responses."""
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def send_json_response(self, code, data):
        """Helper to send JSON responses with CORS headers."""
        response_bytes = json.dumps(data).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", str(len(response_bytes)))
        self.end_headers()
        self.wfile.write(response_bytes)


class DualStackThreadingServer(ThreadingHTTPServer):
    address_family = socket.AF_INET6
    daemon_threads = True
    allow_reuse_address = True

    def server_bind(self):
        try:
            self.socket.setsockopt(socket.IPPROTO_IPV6, socket.IPV6_V6ONLY, 0)
        except Exception:
            pass
        return super().server_bind()


def keep_alive_worker():
    """Background thread that automatically pings the live Render server every 10 minutes to prevent sleep."""
    # Allow 25 seconds for server startup and DNS/network routing
    time.sleep(25)

    # Target URL: Render environment variable or live production URL
    target_domain = os.environ.get("RENDER_EXTERNAL_URL", "").strip()
    if not target_domain or "localhost" in target_domain:
        target_domain = "https://getknowledge.onrender.com"

    ping_url = target_domain.rstrip("/") + "/api/health"
    print(f"[*] Self-keepalive worker started targeting: {ping_url}")

    while True:
        try:
            req = urllib.request.Request(
                ping_url,
                headers={"User-Agent": "GetKnowledge-SelfPing/1.0"}
            )
            with urllib.request.urlopen(req, timeout=12) as resp:
                print(f"[keepalive] Pinged {ping_url} -> Status {resp.status}")
        except Exception as e:
            print(f"[keepalive] Ping note: {e}")

        # Sleep for 10 minutes (600 seconds) - well within Render's 15-minute idle limit
        time.sleep(600)


def main():
    ensure_xlsx_exists()

    # Automatically start the keep-alive background worker on Render
    threading.Thread(target=keep_alive_worker, daemon=True).start()

    server_address = ("::", PORT)
    try:
        httpd = DualStackThreadingServer(server_address, GetKnowledgeRequestHandler)
    except Exception:
        # Fallback to standard IPv4 threading server
        httpd = ThreadingHTTPServer(("", PORT), GetKnowledgeRequestHandler)
        httpd.daemon_threads = True
        httpd.allow_reuse_address = True

    print(f"================================================================")
    print(f"  Get Knowledge Web & XLSX Backend Server Running (Multi-Threaded)")
    print(f"  Local Website URL: http://localhost:{PORT}")
    print(f"  XLSX File Target : {os.path.abspath(XLSX_FILENAME)}")
    print(f"  Official Email   : {OFFICIAL_EMAIL}")
    print(f"  Official Phone   : {OFFICIAL_PHONE}")
    print(f"================================================================")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n[*] Server stopped.")
        httpd.server_close()


if __name__ == "__main__":
    main()
